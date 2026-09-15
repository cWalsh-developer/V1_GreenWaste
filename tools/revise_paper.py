from copy import deepcopy
from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\Users\conno\OneDrive - Kingston University\MSc_Artificial_Intelligence\GreenWaste\AI-Driven Framework for CO₂ Decision Support in Waste Management.docx")
OUTPUT = Path(r"D:\Green Waste\V1_GreenWaste\AI-Driven Framework for CO2 Decision Support - reviewed with citations.docx")

doc = Document(SOURCE)

replacements = {
20: "Bulky household waste creates a practical sustainability problem because decisions made during collection influence whether an item is reused, recycled, incinerated or sent to landfill. Computer vision can recognise waste items, but an image alone cannot provide an exact item-level carbon value because size, mass, material, condition and treatment route remain uncertain. This paper develops a portable decision-support framework that combines RGB-D sensing, object detection, similarity-based product matching and scenario-based life cycle assessment. The system uses an item class and depth-derived size cues to retrieve comparable products and infer plausible material and weight ranges. These ranges are then combined with route-specific emission factors to compare end-of-life scenarios. Instead of presenting a falsely precise value, the framework reports indicative CO₂e ranges and the assumptions behind them. The first version focuses on bulky household items and is intended to test whether this staged, uncertainty-aware approach can support decisions at the point of collection. [Replace this abstract after completing the results, discussion and conclusion: report the final dataset, model, key numerical results, main limitation and conclusion.]",
22: "Waste management is an increasing environmental and operational challenge as urbanisation and consumption place more pressure on collection and treatment systems. The Global Waste Management Outlook 2024 projects that municipal solid waste generation could rise from 2.1 billion tonnes in 2023 to 3.8 billion tonnes by 2050 without urgent action (United Nations Environment Programme, 2024). At collection, decisions about reuse, recycling, incineration and landfill affect both emissions and the materials retained within a circular economy.",
23: "Artificial intelligence can improve the information available at this point. Deep-learning systems can classify and locate waste in images using convolutional neural networks, transfer learning and object detectors such as YOLO (Adedeji and Wang, 2019; Terven, Córdova-Esparza and Romero-González, 2023). However, recognising a chair, table or sofa does not reveal its mass, material composition, condition or suitable end-of-life route.",
24: "This limitation is important for bulky household waste, where items within one class vary considerably. A wooden dining chair, upholstered office chair and metal-framed chair may share a label but differ in mass, materials and reuse potential. Class prediction alone can therefore oversimplify an environmental decision.",
25: "RGB-D sensing adds physical information that a colour image cannot reliably supply. Depth can support object localisation and approximate dimensions when the camera is calibrated and the object is adequately segmented (SrirangamSridharan et al., 2018). It does not measure mass, but it can constrain whether an item is likely to be small, medium or large and improve subsequent reference matching.",
26: "Life cycle assessment (LCA) provides a structured way to compare waste-treatment scenarios, commonly using global warming potential in kilograms of carbon dioxide equivalent (kg CO₂e) (Yadav and Samadder, 2018). Its results depend on the functional unit, system boundary, material data, allocation choices and emission factors (Igos et al., 2019). Camera-derived inputs therefore do not justify an exact item-level result.",
27: "This study develops an AI-assisted framework that links object detection, depth-derived size cues, reference-product matching and scenario-based LCA. The first version focuses on bulky household waste and is designed for portable edge hardware. Its purpose is decision support rather than certified carbon accounting: each stage records its assumptions and passes uncertainty forward instead of hiding it in one fixed value.",
28: "The objectives are to evaluate recognition of selected bulky-waste classes; test whether depth sensing provides useful size estimates; infer plausible material and weight ranges from structured product data; and compare reuse, recycling, incineration and landfill through indicative CO₂e ranges. The main contribution is the integration of these stages into one auditable pipeline.",
33: "Computer vision has been widely applied to waste classification and sorting. Fang et al. (2023) review methods including image classification, segmentation, hyperspectral imaging and robotic perception, while also identifying problems caused by cluttered and variable operating conditions. This supports the use of localisation or segmentation when the target item must be separated from its background.",
34: "White et al. (2020) demonstrated WasteNet on a Jetson Nano and reported 97% test accuracy across six TrashNet categories. This result shows that edge inference is feasible, although performance on a curated classification dataset should not be treated as evidence of the same accuracy on bulky waste in uncontrolled scenes.",
35: "Li and Grammenos (2022) also trained waste classifiers for embedded platforms, supporting on-device inference where cloud access is unreliable. For this project, an object detector is more useful than a whole-image classifier because it supplies both a class and a region that can be aligned with depth data (Terven, Córdova-Esparza and Romero-González, 2023).",
36: "Street-level litter research further shows that changing viewpoints, scale and background clutter make field deployment harder than benchmark classification (Mandhati et al., 2024). Recognition errors would also propagate into material matching and CO₂e estimates, so confidence and failure cases need to be reported.",
39: "RGB-D cameras provide aligned colour and distance measurements that can support object localisation and approximate dimensions in real-world units (SrirangamSridharan et al., 2018). The wider RGB-D literature also shows that results depend on the sensor, calibration, range and scene conditions (Lopes, Souza and Pedrini, 2022; Grenzdörffer, Günther and Hertzberg, 2020).",
40: "Rustler, Volprecht and Hoffmann (2025) found that camera performance changed with distance and task. Their comparison involved the RealSense D435 and D455, ZED 2 and OAK-D Pro, so it supports sensor-conscious testing but does not directly rank the D415 against the ZED 2i used here.",
41: "The D415 is a sensible starting sensor because it provides compact RGB-D capture and has been characterised for close-range measurement and stable tracking under controlled conditions (Moghari et al., 2024). Research using D415 point clouds also shows that geometric features such as height and volume can support later inference, although this requires segmentation and 3D processing rather than treating depth as a direct measurement of mass (Sonar et al., 2024).",
45: "Life cycle assessment evaluates environmental impacts across a defined goal, functional unit and system boundary, followed by inventory analysis, impact assessment and interpretation (International Organization for Standardization, 2006). It is widely used to compare municipal-waste strategies (Nurzhan, Ruan and Chen, 2025). Waste studies must represent alternative treatment routes because the same material can enter several technically different pathways (Towa, Zeller and Achten, 2020).",
46: "LCA is not automatically a complete decision model. Ciacci and Passarini (2020) show that it can be combined with methods such as material-flow analysis, geographic information systems or uncertainty modelling when spatial, resource or operational factors matter. This is relevant here because a camera-based estimate contains uncertainty before the LCA calculation begins.",
49: "LCA results are sensitive to data quality and modelling assumptions, so uncertainty should be reported rather than hidden behind a single value (Igos et al., 2019). Scenario analysis is useful when lifetime, servicing or background-system assumptions could change the conclusion (De Bortoli and Christoforou, 2020).",
50: "Commercial tools such as SimaPro and GaBi can produce comparable results when they use harmonised inventories and assumptions, but tool choice does not remove methodological uncertainty (Herrmann and Moltesen, 2015). Open-source Python packages such as lcpy can also support parametric and uncertainty-aware workflows, although they still require suitable inventory data and impact factors (Gkousis and Katsou, 2025).",
54: "The literature covers waste recognition, RGB-D measurement and LCA separately, but it does not establish a portable workflow that links all three for item-level bulky-waste decisions. This project addresses that practical gap by combining edge-based detection, approximate geometry, reference matching and scenario comparison. Its output is comparative decision support, not an exact carbon footprint.",
56: "The V1 method uses a staged pipeline rather than predicting environmental impact directly from an image. It recognises and locates an item, estimates approximate dimensions from depth, retrieves similar reference products, infers plausible material and mass ranges, and compares end-of-life scenarios. The focus is bulky household waste, where variation within a class makes transparent ranges more defensible than certified or exact values.",
59: "The pipeline is: RGB-D capture → object detection → depth-based size estimation → reference-based material and weight inference → scenario-based CO₂e comparison → decision-support output. Each stage records an estimate and confidence or range so that uncertainty can be carried into the final result.",
67: "The output contains the detected class and confidence, approximate dimensions or size category, matched reference products, likely material family, weight range and modelled CO₂e range for each route. This gives an operator a comparison without presenting one uncertain calculation as fact.",
70: "The prototype uses an Intel RealSense D415, a Stereolabs ZED 2i, an NVIDIA Jetson Orin Nano Developer Kit and portable power. The D415 is the primary sensor for controlled indoor development, supplying aligned colour and depth streams. The ZED 2i is reserved for later, wider-range or field-like comparison rather than being assumed to perform better without testing.",
71: "The Jetson provides a route to local inference where network access may be unavailable. V1 tests whether this portable setup can support the workflow; it does not claim a deployment-ready vehicle system. Hardware suitability will need to be evaluated through accuracy, latency, power and environmental testing.",
76: "Two data sources are used: captured RGB-D observations and a structured reference library containing product category, dimensions, materials and weight. Each capture stores the RGB image, depth map, camera model, resolution, depth scale and camera intrinsics. Intrinsics and depth scale are required to convert image coordinates and raw depth into approximate physical units.",
77: "Items are captured under controlled lighting with limited background clutter. Camera distance and viewpoint should be standardised, and the sensor should be warmed up and calibrated according to its operating procedure. Manual width, height, depth and mass measurements are recorded as ground truth so that geometric and reference-based estimates can be evaluated rather than only demonstrated.",
78: "The product library is used for similarity-based inference, not exact product identification. A detected chair may be compared with IKEA-style chair records to constrain likely dimensions, materials and weight, but the system should not claim it is a particular product.",
82: "V1 limits detection to a small set of bulky household categories so the full pipeline can be evaluated within the project timeframe. The detector returns an item class, confidence score and bounding box or segmentation mask. Localisation is needed because only the corresponding depth region should be analysed.",
87: "A detector is therefore more appropriate than whole-image classification. However, a rectangular bounding box may include background pixels, so a segmentation mask should be used where available or the depth region should be filtered robustly. The recognition stage does not infer exact materials; it supplies the class and location used by later stages.",
91: "For each detected region, valid depth pixels are filtered to remove missing values and obvious background points. A robust central depth, such as the median, estimates object distance. With aligned depth and camera intrinsics, image coordinates can be deprojected into 3D points and used to estimate visible width and height. A single frontal view does not reliably recover full object depth or dimensions hidden by occlusion.",
92: "V1 therefore targets approximate dimensions rather than full 3D reconstruction. Estimates can be grouped into broad size bands for reference matching, but they must first be compared with manual measurements using an error measure such as mean absolute error or mean absolute percentage error.",
93: "Depth addresses the scale ambiguity of a 2D image, but its accuracy is affected by range, angle, reflective or textureless surfaces, occlusion, calibration and segmentation. These conditions should be recorded as limitations or rejected through quality thresholds rather than silently passed to the next stage.",
97: "Reference matching begins by filtering products to the detected class and normalising their dimensions to a consistent unit and orientation. Candidate similarity can then be calculated from dimension differences, with the nearest records used to form a plausible weight range and summarise material descriptions.",
98: "Material descriptions are mapped to broad families such as wood-dominant, metal-dominant, plastic-dominant, textile/foam-dominant or mixed. Weight is reported as a range across the best matches rather than copied from one product. The number of neighbours and similarity threshold must be fixed before evaluation and reported so that the process is repeatable.",
99: "These matches are proxies, not measurements. Products with similar outer dimensions can differ in construction and mass, and the reference library may not represent the collected-waste population. Validation should therefore compare inferred ranges with known test-item mass and report coverage: the proportion of true masses that fall inside the predicted range.",
103: "The final stage combines the estimated mass range with material- and route-specific factors for reuse, closed-loop recycling, incineration with energy recovery and landfill. The calculation is an end-of-life scenario model rather than a complete cradle-to-grave LCA. It excludes the item's original manufacture and use, except for the avoided-production proxy used in the reuse scenario.",
106: "As a worked example, the two chair composition profiles produce landfill intensities of approximately 0.287 and 0.698 kg CO₂e/kg. If reference matching estimates a mass of 5–8 kg, the four endpoint products give an overall landfill range of approximately 1.43–5.59 kg CO₂e. The width of this range comes from uncertainty in both mass and assumed composition; it is not a confidence interval derived from repeated observations.",
107: "Reuse needs separate treatment because any credit depends on whether reuse actually displaces production of a replacement item. Condition, demand, refurbishment and transport therefore affect the scenario. Recycling and energy-recovery credits also depend on substitution assumptions, while landfill depends on material degradation and gas management.",
108: "Routes should be compared only on the same basis and subject to practical constraints such as item condition, contamination, safety, local facility availability and transport. Choosing the lowest upper bound is a conservative option, but it is not automatically the best ranking rule; V1 should report interval overlap and flag cases where uncertainty prevents a clear recommendation.",
109: "The final display should show the detected item, confidence, size estimate, reference-derived mass and material ranges, the factor source and boundary, each route’s CO₂e interval, and any quality warning. This makes the result auditable and keeps the system within its role as decision support rather than formal carbon accounting.",
112: "Recommended route: report only when the modelled comparison and operational rules support a clear choice; otherwise report that the result is uncertain.",
128: "United Nations Environment Programme (2024) Global Waste Management Outlook 2024: Beyond an age of waste – Turning rubbish into a resource. Available at: https://www.unep.org/resources/global-waste-management-outlook-2024 (Accessed: 20 August 2026).",
131: "Gkousis, S. and Katsou, E. (2025) ‘lcpy: an open-source Python package for parametric and dynamic Life Cycle Assessment and Life Cycle Costing’. arXiv. Available at: https://doi.org/10.48550/arXiv.2506.13744.",
134: "Igos, E. et al. (2019) ‘How to treat uncertainties in life cycle assessment studies?’, The International Journal of Life Cycle Assessment, 24(4), pp. 794–807. Available at: https://doi.org/10.1007/s11367-018-1477-1.",
136: "Li, X. and Grammenos, R. (2022) ‘A Smart Recycling Bin Using Waste Image Classification at the Edge’. arXiv. Available at: https://doi.org/10.48550/arXiv.2210.00448.",
147: "White, G. et al. (2020) ‘WasteNet: Waste Classification at the Edge for Smart Bins’. arXiv. Available at: https://doi.org/10.48550/arXiv.2006.05873.",
143: "SrirangamSridharan, S. et al. (2018) ‘Object Localization and Size Estimation from RGB-D Images’. arXiv. Available at: https://doi.org/10.48550/arXiv.1808.00641.",
}

remove = {29, 30, 37, 42, 43, 47, 51, 52, 57, 60, 61, 62, 63, 64, 65, 66, 68, 72, 73, 74, 79, 80, 83, 84, 85, 86, 88, 89, 94, 95, 100, 101, 104, 110, 113, 135, 144}

for index, text in replacements.items():
    doc.paragraphs[index].text = text

for index in sorted(remove, reverse=True):
    paragraph = doc.paragraphs[index]
    paragraph._element.getparent().remove(paragraph._element)


def insert_after(paragraph, text, style=None):
    new_paragraph = deepcopy(paragraph._element)
    for child in list(new_paragraph):
        new_paragraph.remove(child)
    paragraph._element.addnext(new_paragraph)
    inserted = paragraph._parent.add_paragraph()
    inserted._element.getparent().remove(inserted._element)
    inserted._element = new_paragraph
    inserted._p = new_paragraph
    if style:
        inserted.style = style
    inserted.add_run(text)
    return inserted


lca_start = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text.startswith("The final stage combines")
)
factor_source = insert_after(
    lca_start,
    "The material and waste-treatment factors come from the UK Government GHG Conversion Factors for Company Reporting 2025 (Department for Energy Security and Net Zero, 2025a). The project uses the flat file intended for automated processing. Published values in kg CO₂e per tonne are divided by 1,000 to give kg CO₂e per kg. Wood, metal, average plastic, glass and clothing categories are mapped to the broad material families used by the system. Foam uses average-plastic factors, while furniture upholstery uses clothing factors, because suitable furniture-specific entries are unavailable. These substitutions are project proxies and are marked as lower-quality assumptions.",
    lca_start.style,
)
composition = insert_after(
    factor_source,
    "The camera does not measure material composition. Each detected class therefore has two alternative proxy composition profiles. For example, a chair may be represented by different fractions of wood, metal, textile and foam. For material fractions pₘ and route-specific factors EFₘ,ᵣ, the composite intensity for route r is EFᵣ = Σₘ(pₘ × EFₘ,ᵣ). The profiles produce alternative intensities, and the outer bounds are carried forward. These profiles are transparent project assumptions derived for uncertainty testing; they are not supplied by the UK Government dataset or confirmed compositions of the photographed items.",
    lca_start.style,
)
calculation = insert_after(
    composition,
    "For an inferred mass interval [Mₗ, Mₕ] and intensity interval [EFₗ, EFₕ], the reported CO₂e interval is the minimum and maximum of MₗEFₗ, MₗEFₕ, MₕEFₗ and MₕEFₕ. Checking all four products is necessary because reuse is expressed as a negative avoided-impact proxy. Recycling, incineration and landfill use waste-disposal factors and do not automatically include avoided virgin-material production. Substitution credits depend on how much secondary material or continued product use genuinely displaces primary production (European Commission Joint Research Centre, 2010). Reuse is therefore conditional on item condition, demand, transport and actual substitution.",
    lca_start.style,
)
decision = insert_after(
    calculation,
    "The displayed CO₂e ranges remain the unadjusted factor results. Route recommendation is a separate decision rule: when condition is unknown or the item is not reusable, the code ranks disposal routes using the upper CO₂e bound plus a small project-defined waste-hierarchy adjustment of 0 kg CO₂e/kg for recycling, 0.05 for incineration and 0.10 for landfill. This adjustment is a transparent ranking heuristic, not part of the UK factors or the calculated LCA result.",
    lca_start.style,
)
caveat = insert_after(
    decision,
    "This source has an important limitation. Its methodology states that the material-use factors describe procured materials and that, except for landfill, the waste-disposal factors mainly cover collection and delivery to the treatment or disposal point. It warns that these figures do not represent the full environmental impact of alternative waste-management options (Department for Energy Security and Net Zero, 2025b). The current outputs must therefore be described as preliminary route indicators, not comparative LCA results. A stronger implementation should replace or validate these factors using a waste-scenario model such as CarbonWARM2, which was developed for comparing the greenhouse-gas effects of recycling, reuse and disposal choices (WRAP, 2025).",
    lca_start.style,
)

reference_anchor = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text.startswith("De Bortoli, A. and Christoforou")
)
insert_after(
    reference_anchor,
    "Department for Energy Security and Net Zero (2025a) Greenhouse gas reporting: conversion factors 2025. Available at: https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2025 (Accessed: 26 August 2026).",
    "Bibliography",
)
desnz_a = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text.startswith("Department for Energy Security and Net Zero (2025a)")
)
insert_after(
    desnz_a,
    "Department for Energy Security and Net Zero (2025b) 2025 Government greenhouse gas conversion factors for company reporting: Methodology paper. Available at: https://assets.publishing.service.gov.uk/media/6846b0870392ed9b784c0187/2025-GHG-CF-methodology-paper.pdf (Accessed: 26 August 2026).",
    "Bibliography",
)
igos_anchor = next(
    paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("Igos, E.")
)
iso_ref = insert_after(
    igos_anchor,
    "International Organization for Standardization (2006) ISO 14040:2006 Environmental management — Life cycle assessment — Principles and framework. Geneva: ISO. Available at: https://www.iso.org/standard/37456.html (Accessed: 26 August 2026).",
    "Bibliography",
)
eitel_anchor = next(
    paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("Eitel, A.")
)
insert_after(
    eitel_anchor,
    "European Commission Joint Research Centre (2010) International Reference Life Cycle Data System (ILCD) Handbook: General guide for Life Cycle Assessment — Detailed guidance. Luxembourg: Publications Office of the European Union. Available at: https://doi.org/10.2788/38479.",
    "Bibliography",
)
white_anchor = next(
    paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("White, G.")
)
insert_after(
    white_anchor,
    "WRAP (2025) Carbon Waste and Resources Metric (CarbonWARM2). Available at: https://www.wrap.ngo/resources/report/carbon-waste-and-resources-metric-carbonwarm2 (Accessed: 26 August 2026).",
    "Bibliography",
)

doc.save(OUTPUT)
print(OUTPUT)
