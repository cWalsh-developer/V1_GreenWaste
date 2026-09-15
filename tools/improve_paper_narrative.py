from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\conno\OneDrive - Kingston University\MSc_Artificial_Intelligence\GreenWaste\AI-Driven Framework for CO2 Decision Support.docx")
BACKUP = SOURCE.with_name("AI-Driven Framework for CO2 Decision Support.before-narrative-edit.docx")
OUTPUT = SOURCE.with_name("AI-Driven Framework for CO2 Decision Support - narrative revision.docx")


TRANSITIONS = {
    "The objectives are to evaluate recognition of selected bulky-waste classes;":
        " These stages also provide the structure for the literature review, which moves from recognition to physical estimation and finally to environmental comparison.",
    "Street-level litter research further shows that changing viewpoints":
        " Recognition alone cannot resolve the physical scale of an item, making depth sensing the next requirement in the proposed pipeline.",
    "Rustler, Volprecht and Hoffmann (2025) compared the D435":
        " These depth-derived measurements are not environmental outcomes themselves, but they provide the physical evidence needed to constrain the later LCA inputs.",
    "Commercial tools such as SimaPro and GaBi":
        " The need to preserve this uncertainty across recognition, measurement and LCA leads directly to the research gap addressed here.",
    "The output contains the detected class and confidence":
        " Implementing this sequence requires hardware that can capture aligned visual and depth information locally.",
    "The Jetson provides a route to local inference":
        " This hardware configuration determines the two forms of data used by the pipeline.",
    "The product library is used for similarity-based inference":
        " Before this library can be searched, the captured item must first be classified and localised.",
    "Depth addresses the scale ambiguity of a 2D image":
        " Once a depth estimate passes these checks, its size cues can be used to search for comparable products.",
    "These matches are proxies, not measurements.":
        " The resulting material and mass ranges then form the uncertain inputs to the scenario calculation.",
    "Recommended route: report only when":
        " These rules define the outputs evaluated in the following Results section.",
}


DISCUSSION = {
    "5.1 Interpretation of Findings": [
        "The results show that the individual stages can be connected into a working decision-support pipeline, but detector performance remains the main constraint. The original five-class cross-validation score suggested that the model had learnt the available dataset well. However, the lower success rate on RealSense captures shows that internal evaluation did not represent the intended operating environment. This matters because a missed or incorrect detection prevents reliable size estimation, reference matching and CO₂e calculation, regardless of how well the later stages function.",
        "Targeted manual correction was particularly valuable for the weak tables_desks class. The reviewed model provided the best balance for the pipeline, while adding pseudo-labels improved stricter localisation but reduced recall. Pseudo-labelling was therefore useful for adding clear examples but did not replace human review of difficult images. The findings support a human-assisted approach in which uncertain cases are rejected or reviewed rather than automatically passed through the system.",
    ],
    "5.2 Limitations": [
        "The main limitation is the size and representativeness of the image dataset. Much of the training data consisted of product-style images, whereas collected waste may appear damaged, partly hidden, poorly lit or surrounded by clutter. Only seven saved RealSense captures were used for the final test, so the one successful route demonstrates technical operation but cannot provide a reliable estimate of field accuracy. The corrected evaluation boxes also began as model predictions, creating a risk of anchoring bias.",
        "Physical measurements remain approximate because they depend on the bounding box, valid depth pixels and a single viewpoint. The system also infers mass and material composition from similar products rather than measuring them. Items with comparable dimensions may differ in construction and weight, so these ranges are not confirmed properties.",
        "The CO₂e output has a further boundary limitation. The UK Government waste-disposal factors mainly represent collection and delivery to the first processing point under Scope 3 accounting. They do not provide a complete comparison of recycling and incineration, explaining why those routes received identical raw values here. Reuse is also conditional on actually displacing another purchase. The calculations are therefore route indicators rather than a comparative LCA or certified carbon footprint.",
        "Finally, prolonged field performance, processing speed, power use, weather and vehicle integration were not evaluated.",
    ],
    "5.3 Practical Implications": [
        "The prototype is most suitable as a triage tool rather than an autonomous decision maker. It could provide a likely category, approximate size and initial route comparison while displaying uncertainty and requesting a manual condition check. Rejecting a low-confidence detection is preferable to producing a confident-looking result from unreliable inputs.",
        "The results also establish the priority for further development: realistic labelled RGB-D images are needed before refining later calculations. Subsequent work should validate dimensions and inferred weight ranges against measurements and replace the current disposal factors with a consistent waste-scenario model.",
    ],
    "6. Conclusions and Future Work": [
        "This project developed a proof-of-concept pipeline linking furniture detection, depth-based size estimation, reference-product matching and scenario-based CO₂e calculation. The original five-class model achieved a mean mAP50 of 0.943 in five-fold cross-validation. In the separate tables_desks evaluation, targeted review improved mAP50 from 0.256 for the original detector to 0.943 for the specialist model; these identical final values came from different test designs and are not directly comparable. The complete pipeline processed one of seven saved RealSense captures, confirming technical integration but limited transfer to natural capture conditions.",
        "The project demonstrates that camera-derived information can support an auditable sequence of indicative waste-routing decisions, but not an exact item-level carbon assessment. Its main contribution is carrying uncertain inputs forward as ranges rather than presenting them as measured facts. Future work should prioritise realistic RGB-D training data, independent field testing, measured size and mass validation, and treatment factors with consistent life-cycle boundaries. Until then, the system should be described as human-assisted decision support rather than a deployment-ready waste-management or carbon-accounting tool.",
    ],
}


def insert_after(paragraph: Paragraph, texts: list[str]) -> None:
    anchor = paragraph._p
    for text in texts:
        new_p = OxmlElement("w:p")
        anchor.addnext(new_p)
        new_paragraph = Paragraph(new_p, paragraph._parent)
        new_paragraph.style = "Normal"
        new_paragraph.add_run(text)
        anchor = new_p


def main() -> None:
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    document = Document(SOURCE)
    added = 0
    for paragraph in document.paragraphs:
        for start, transition in TRANSITIONS.items():
            if paragraph.text.startswith(start) and transition.strip() not in paragraph.text:
                paragraph.add_run(transition)
                added += 1

    paragraphs = list(document.paragraphs)
    for heading, texts in DISCUSSION.items():
        target = next((p for p in paragraphs if p.text.strip() == heading), None)
        if target is None:
            raise RuntimeError(f"Heading not found: {heading}")
        next_text = ""
        index = paragraphs.index(target)
        if index + 1 < len(paragraphs):
            next_text = paragraphs[index + 1].text.strip()
        if not next_text or paragraphs[index + 1].style.name.startswith("Heading"):
            insert_after(target, texts)

    try:
        document.save(SOURCE)
        saved_to = SOURCE
    except PermissionError:
        document.save(OUTPUT)
        saved_to = OUTPUT

    print(f"Saved: {saved_to}")
    print(f"Backup: {BACKUP}")
    print(f"Transitions added: {added}")


if __name__ == "__main__":
    main()
