from pathlib import Path
from docx import Document
import sys


def extract(source: str, output: str) -> None:
    doc = Document(source)
    lines = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else ""
            lines.append(f"P{index:04d}\t[{style}]\t{text}")
    for table_index, table in enumerate(doc.tables):
        lines.append(f"TABLE {table_index}")
        for row in table.rows:
            lines.append("\t".join(cell.text.replace("\n", " | ").strip() for cell in row.cells))
    Path(output).write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    extract(sys.argv[1], sys.argv[2])
