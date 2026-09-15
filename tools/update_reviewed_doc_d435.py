from pathlib import Path
import shutil

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "AI-Driven Framework for CO2 Decision Support - reviewed with citations.docx"
BACKUP = ROOT / "AI-Driven Framework for CO2 Decision Support - reviewed with citations.before-D435-correction.docx"
OUTPUT = ROOT / "AI-Driven Framework for CO2 Decision Support - reviewed with citations - D435 corrected.docx"


REPLACEMENTS = {
    "Rustler, Volprecht and Hoffmann (2025) found that camera performance changed with distance and task. Their comparison involved the RealSense D435 and D455, ZED 2 and OAK-D Pro, so it supports sensor-conscious testing but does not directly rank the D415 against the ZED 2i used here.":
        "Research using the Intel RealSense D435 shows that depth reliability depends on the capture conditions. Condotta et al. (2020) evaluated the D435 at different distances and image positions in indoor and outdoor settings. Their results showed that measurement variability increased with distance, supporting the need to control the camera position and treat the dimensions produced by the prototype as estimates rather than exact measurements.",
    "The D415 is a sensible starting sensor because it provides compact RGB-D capture and has been characterised for close-range measurement and stable tracking under controlled conditions (Moghari et al., 2024). Research using D415 point clouds also shows that geometric features such as height and volume can support later inference, although this requires segmentation and 3D processing rather than treating depth as a direct measurement of mass (Sonar et al., 2024).":
        "Rustler, Volprecht and Hoffmann (2025) compared the D435 with the RealSense D455, ZED 2 and OAK-D Pro using planar surfaces, complex shapes and household objects. They found that the D435 was particularly suitable for short-range object perception at approximately 90–100 cm, while its accuracy and precision reduced at longer distances. This is relevant to the present system because the D435 is used at relatively close range to locate furniture and obtain approximate depth measurements.",
    "The prototype uses an Intel RealSense D415, a Stereolabs ZED 2i, an NVIDIA Jetson Orin Nano Developer Kit and portable power. The D415 is the primary sensor for controlled indoor development, supplying aligned colour and depth streams. The ZED 2i is reserved for later, wider-range or field-like comparison rather than being assumed to perform better without testing.":
        "The prototype uses an Intel RealSense D435 RGB-D camera, an NVIDIA Jetson Orin Nano Developer Kit and portable power. The D435 supplies the colour frame used for object detection and an aligned depth frame used to estimate the distance and approximate dimensions of the detected item. This configuration is intended to test a portable processing route rather than claim a deployment-ready collection system.",
}


REFERENCES = [
    "Condotta, I.C.F.S., Brown-Brandl, T.M., Pitla, S.K., Stinn, J.P. and Silva-Miranda, K.O. (2020) ‘Evaluation of low-cost depth cameras for agricultural applications’, Computers and Electronics in Agriculture, 173, 105394. Available at: https://doi.org/10.1016/j.compag.2020.105394.",
    "Rustler, L., Volprecht, V. and Hoffmann, M. (2025) ‘Empirical Comparison of Four Stereoscopic Depth Sensing Cameras for Robotics Applications’, IEEE Access. Available at: https://doi.org/10.1109/ACCESS.2025.3560810.",
]


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    document = Document(SOURCE)
    replaced = set()
    for paragraph in document.paragraphs:
        current = paragraph.text.strip()
        if current in REPLACEMENTS:
            paragraph.text = REPLACEMENTS[current]
            replaced.add(current)

    missing = set(REPLACEMENTS) - replaced
    if missing:
        raise RuntimeError(f"Expected paragraphs not found: {len(missing)}")

    existing_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for reference in REFERENCES:
        if reference not in existing_text:
            document.add_paragraph(reference)

    try:
        document.save(SOURCE)
        saved_to = SOURCE
    except PermissionError:
        document.save(OUTPUT)
        saved_to = OUTPUT
    print(f"Updated: {saved_to}")
    print(f"Backup:  {BACKUP}")
    print(f"Replaced paragraphs: {len(replaced)}")


if __name__ == "__main__":
    main()
