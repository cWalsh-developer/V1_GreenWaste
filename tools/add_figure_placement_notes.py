from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor


SOURCE = Path(
    r"D:\Green Waste\V1_GreenWaste\AI-Driven Framework for CO2 Decision Support - reviewed with citations.docx"
)
OUTPUT = Path(
    r"D:\Green Waste\V1_GreenWaste\AI-Driven Framework for CO2 Decision Support - figure placement notes.docx"
)


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = "Normal"
    run = inserted.add_run(text)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return inserted


doc = Document(SOURCE)

notes = {
    "4.1. Prototype Pipeline Output": [
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.1] After the paragraph reporting the successful chair detection "
            "(confidence 0.837 and estimated dimensions 149.3 × 112.5 × 59.0 cm), insert "
            "docs/results_assets/capture_20260527_022745_annotated_detection.png. Lead in with: “Figure 4.1 "
            "shows the RGB detection and corresponding depth-based size output for a successful RealSense chair "
            "capture.” Caption: “Figure 4.1. RealSense chair capture processed by the prototype. The item was "
            "detected as chair_seating with 0.837 confidence, and the detected region was passed to the "
            "depth-based size-estimation stage.”"
        ),
    ],
    "4.2. Item Recognition and Size Estimation Results": [
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.2] After the table comparing the original and IKEA-plus-curated "
            "five-fold cross-validation metrics, insert docs/results_assets/cv_original_vs_curated_comparison.png. "
            "Lead in with: “Figure 4.2 visualises the reduction in detector performance after curated "
            "natural-background images were added.” Caption: “Figure 4.2. Mean detector performance for the "
            "original and IKEA-plus-curated five-fold cross-validation experiments. Error bars represent the "
            "standard deviation across folds.”"
        ),
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.3] After discussing the fold-level mAP50–95 range and variation, "
            "insert docs/results_assets/cv_fold_map50_95_comparison.png. Lead in with: “The fold-level results in "
            "Figure 4.3 show that the reduction occurred throughout the curated experiment rather than being caused "
            "by one poor split.” Caption: “Figure 4.3. Fold-level mAP50–95 for the original and "
            "IKEA-plus-curated datasets. The original dataset produced stronger localisation performance across "
            "all five folds.”"
        ),
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.4] Immediately after the main experiment-comparison table, insert "
            "docs/results_assets/experiment_map50_95_summary.png. Lead in with: “Figure 4.4 places the "
            "cross-validation results alongside the final fitting and weak-label experiments.” Caption: “Figure "
            "4.4. mAP50–95 across the main detector experiments. Train-as-validation results are included for "
            "transparency but are not treated as unbiased estimates of generalisation.”"
        ),
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.5] After the paragraph explaining confusion between chairs, sofas, "
            "storage units and tables, insert docs/results_assets/curated_stratified_confusion_matrix_normalized.png. "
            "Lead in with: “The normalised confusion matrix in Figure 4.5 provides a class-level view of these "
            "errors.” Caption: “Figure 4.5. Normalised confusion matrix for the IKEA-plus-curated stratified "
            "experiment, showing remaining class confusion and background-related detection errors.”"
        ),
        (
            "[OPTIONAL APPENDIX NOTE] Put original_cv_fold1_confusion_matrix_normalized.png, "
            "original_cv_fold1_pr_curve.png and the training-curve figures in an appendix unless one is needed to "
            "support a specific point. Avoid placing every training diagnostic in the main Results section."
        ),
    ],
    "4.3. Reference Matching and CO₂e Scenario Outputs": [
        (
            "[FIGURE PLACEMENT NOTE — FIGURE 4.6] After the storage-capture scenario-results table, insert "
            "docs/results_assets/capture_20260618_222719_lca_scenarios.png. Lead in with: “Figure 4.6 visualises "
            "the preliminary CO₂e route indicators generated for the storage capture.” Caption: “Figure 4.6. "
            "Preliminary CO₂e route indicators for the storage capture. Negative reuse values represent a "
            "conditional avoided-production proxy. The other values use UK Government waste-disposal factors and "
            "do not constitute a complete comparative LCA.”"
        ),
        (
            "[OPTIONAL APPENDIX NOTE] Place capture_20260618_232017_lca_scenarios.png and additional annotated "
            "RealSense examples in an appendix. Figure 4.6 already demonstrates the scenario-output format in the "
            "main text."
        ),
    ],
}

for heading, section_notes in notes.items():
    anchor = next(p for p in doc.paragraphs if p.text.strip() == heading)
    for note in reversed(section_notes):
        insert_after(anchor, note)

doc.save(OUTPUT)
print(OUTPUT)
