from pathlib import Path

from docx import Document


SOURCE = Path(r"D:\Green Waste\V1_GreenWaste\current_report_for_output_update.docx")
OUTPUT = Path(r"D:\Green Waste\V1_GreenWaste\current_report_with_output_summary.docx")


ROWS = [
    ("Output field", "Final pipeline output"),
    ("Detected item", "Beds and mattresses"),
    ("Detection confidence", "0.363"),
    ("Size category", "Medium"),
    ("Estimated dimensions", "101.1 × 47.9 × 44.7 cm"),
    ("Likely material family", "Mixed (proxy estimate)"),
    ("Estimated weight range", "29.97–46.18 kg"),
    ("Reuse", "−590.885 to −289.330 kg CO₂e (conditional)"),
    ("Closed-loop recycling", "0.118–0.182 kg CO₂e"),
    ("Incineration with energy recovery", "0.118–0.182 kg CO₂e"),
    ("Landfill", "7.533–12.459 kg CO₂e"),
    ("Recommended route", "Recycle while condition is unknown; reuse if confirmed suitable"),
    ("Quality warning", "Preliminary Scope 3 route indicators, not a complete comparative LCA"),
]


def main() -> None:
    document = Document(SOURCE)
    target_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Reference matching estimated a weight range")
    )
    target_paragraph.text = (
        "The successful capture produced the consolidated decision-support output "
        "shown below, using the format defined in Section 3."
    )

    old_table = next(
        table
        for table in document.tables
        if table.cell(0, 0).text.strip() == "Scenario"
    )

    new_table = document.add_table(rows=len(ROWS), cols=2)
    new_table.style = "Table Grid"
    for row_index, (field, value) in enumerate(ROWS):
        new_table.cell(row_index, 0).text = field
        new_table.cell(row_index, 1).text = value
        if row_index == 0:
            for cell in new_table.rows[row_index].cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True

    old_table._tbl.addprevious(new_table._tbl)
    old_table._element.getparent().remove(old_table._element)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
